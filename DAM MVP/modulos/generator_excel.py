import os
from datetime import datetime
import openpyxl


MAPA_CELDAS = {

#Celdas

    ("factura", "numero_factura"): "C4",
    ("factura", "fecha_factura"): "C5",
    ("factura", "importador"): "C6",
    ("factura", "ruc"): "C7",
    ("factura", "exportador"): "C8",
    ("factura", "pais_origen"): "C9",
    ("factura", "valor_total"): "C10",
    ("factura", "valor_fob"): "C10",
    ("factura", "moneda"): "C11",
    ("factura", "incoterm"): "C12",

#Packing List

    ("packing_list", "numero_bultos"): "C15",
    ("packing_list", "tipo_bultos"): "C16",
    ("packing_list", "numero_pallets"): "C17",
    ("packing_list", "peso_neto"): "C18",
    ("packing_list", "peso_bruto"): "C19",
    ("packing_list", "volumen"): "C20",
    ("packing_list", "numero_contenedor"): "C21",

#BL

    ("bill_of_lading", "numero_bl"): "C24",
    ("bill_of_lading", "shipper"): "C25",
    ("bill_of_lading", "consignatario"): "C26",
    ("bill_of_lading", "notify_party"): "C27",
    ("bill_of_lading", "numero_contenedor"): "C28",
    ("bill_of_lading", "numero_precinto"): "C29",
    ("bill_of_lading", "buque"): "C30",
    ("bill_of_lading", "viaje"): "C31",
    ("bill_of_lading", "puerto_origen"): "C32",
    ("bill_of_lading", "puerto_carga"): "C32",
    ("bill_of_lading", "puerto_destino"): "C33",
    ("bill_of_lading", "puerto_descarga"): "C33",
    ("bill_of_lading", "peso_bruto"): "C34",
    ("bill_of_lading", "descripcion_mercancia"): "C35",

#certificado

    ("certificado_seguro", "numero_certificado"): "C38",
    ("certificado_seguro", "numero_poliza"): "C39",
    ("certificado_seguro", "asegurado"): "C40",
    ("certificado_seguro", "numero_factura"): "C41",
    ("certificado_seguro", "numero_bl"): "C42",
    ("certificado_seguro", "suma_asegurada"): "C43",
    ("certificado_seguro", "moneda"): "C44",
    ("certificado_seguro", "prima_seguro"): "C45",
    ("certificado_seguro", "buque"): "C46",
    ("certificado_seguro", "origen"): "C47",
    ("certificado_seguro", "destino"): "C48",
}


#Genera el excel

def generar_excel(
    datos_caso: dict,
    ruta_plantilla: str,
    carpeta_salida: str
) -> str:

    if not os.path.exists(ruta_plantilla):
        crear_plantilla_ejemplo(ruta_plantilla)

    os.makedirs(carpeta_salida, exist_ok=True)

    libro = openpyxl.load_workbook(ruta_plantilla)

    hoja = libro["DAM"]

    documentos = datos_caso.get("documentos", {})

    for (tipo_doc, campo), celda in MAPA_CELDAS.items():

        documento = documentos.get(tipo_doc)

        # Si ese documento no existe, simplemente no hacemos nada
        if not documento:
            continue

        datos_documento = documento.get("datos", {})

        valor = datos_documento.get(campo)

        # Si el dato existe, escribirlo
        if valor is not None and valor != "":
            hoja[celda] = valor

#Trazabilidad

    hoja["C1"] = datos_caso.get(
        "nombre_caso",
        ""
    )

    hoja["C2"] = datetime.now().strftime(
        "%Y-%m-%d %H:%M"
    )

#save

    id_caso = datos_caso.get(
        "id_caso",
        "CASO"
    )

    nombre_salida = f"DAM_{id_caso}.xlsx"

    ruta_salida = os.path.join(
        carpeta_salida,
        nombre_salida
    )

    libro.save(ruta_salida)

    return ruta_salida

#Crear la plntlla

def crear_plantilla_ejemplo(ruta_plantilla: str) -> None:

    os.makedirs(
        os.path.dirname(ruta_plantilla),
        exist_ok=True
    )

    # Si ya existe la plantilla, abrirla.
    # Si no existe, crear una nueva.
    if os.path.exists(ruta_plantilla):

        libro = openpyxl.load_workbook(
            ruta_plantilla
        )

        if "DAM" in libro.sheetnames:
            hoja = libro["DAM"]
        else:
            hoja = libro.active
            hoja.title = "DAM"

    else:

        libro = openpyxl.Workbook()

        hoja = libro.active

        hoja.title = "DAM"


    filas = [

#GENERAL

        ("A1", "Caso:"),
        ("A2", "Fecha de generación:"),

#FACTURA

        ("A3", "FACTURA COMERCIAL"),
        ("A4", "N° Factura:"),
        ("A5", "Fecha factura:"),
        ("A6", "Importador:"),
        ("A7", "RUC:"),
        ("A8", "Exportador:"),
        ("A9", "País de origen:"),
        ("A10", "Valor total / FOB:"),
        ("A11", "Moneda:"),
        ("A12", "Incoterm:"),

#PACKING LIST

        ("A14", "PACKING LIST"),
        ("A15", "N° Bultos:"),
        ("A16", "Tipo de bultos:"),
        ("A17", "N° Pallets:"),
        ("A18", "Peso neto:"),
        ("A19", "Peso bruto:"),
        ("A20", "Volumen:"),
        ("A21", "N° Contenedor:"),

#BL

        ("A23", "BILL OF LADING"),
        ("A24", "N° B/L:"),
        ("A25", "Shipper / Remitente:"),
        ("A26", "Consignatario:"),
        ("A27", "Notify Party:"),
        ("A28", "N° Contenedor:"),
        ("A29", "N° Precinto:"),
        ("A30", "Buque:"),
        ("A31", "Viaje:"),
        ("A32", "Puerto de carga:"),
        ("A33", "Puerto de descarga:"),
        ("A34", "Peso bruto:"),
        ("A35", "Descripción mercancía:"),

#CERTIFICADO DE SEGURO

        ("A37", "CERTIFICADO DE SEGURO"),
        ("A38", "N° Certificado:"),
        ("A39", "N° Póliza:"),
        ("A40", "Asegurado:"),
        ("A41", "Factura relacionada:"),
        ("A42", "B/L relacionado:"),
        ("A43", "Suma asegurada:"),
        ("A44", "Moneda:"),
        ("A45", "Prima de seguro:"),
        ("A46", "Buque:"),
        ("A47", "Origen:"),
        ("A48", "Destino:"),
    ]

    for celda, texto in filas:
        hoja[celda] = texto

    libro.save(ruta_plantilla)